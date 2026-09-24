"""Synthetic batch tests: no corpus reads and no external embedding requests."""

import importlib.util
import json
import sys
import tempfile
import unittest
from itertools import pairwise
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

SPEC = importlib.util.spec_from_file_location(
    "bulk_window_corpus", Path(__file__).parents[1] / "bulk_window_corpus.py"
)
app = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = app
SPEC.loader.exec_module(app)


class BatchTests(unittest.TestCase):
    def test_input_selection_uses_exact_hash_even_for_misnamed_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "misnamed.docx"
            source.write_bytes(b"old Word container")
            converted = Path(tmp) / "converted.docx"
            converted.write_bytes(b"converted OOXML")
            row = {
                "source_path": str(source),
                "source_sha256": app.sha(source),
                "converted_path": str(converted),
                "converted_sha256": app.sha(converted),
            }
            wrong_version = {**row, "converted_sha256": "not-expected"}
            inp, provenance = app.select_input(
                source, app.sha(source), app.sha(converted), [row, wrong_version]
            )
            self.assertEqual(inp, converted)
            self.assertEqual(provenance, row)
            with self.assertRaises(ValueError):
                app.select_input(source, app.sha(source), "unknown", [row])
            converted.write_bytes(b"changed")
            with self.assertRaises(ValueError):
                app.select_input(source, app.sha(source), row["converted_sha256"], [row])

    def test_windows_reconstruct_and_bind_embedding(self):
        text = "中文🙂XYZ。" * 401
        rows = app.windows("doc-a", "合成法", text)
        rebuilt = rows[0]["text"]
        for previous, current in pairwise(rows):
            self.assertEqual(previous["end_char"] - current["start_char"], 200)
            rebuilt += current["text"][200:]
        self.assertEqual(rebuilt, text)
        for c in rows:
            self.assertEqual(c["text"], text[c["start_char"] : c["end_char"]])
            self.assertEqual(
                c["embedding_text"], "法律名称：合成法\n正文：" + c["text"]
            )
        self.assertNotEqual(
            rows[0]["chunk_id"], app.windows("doc-b", "合成法", text)[0]["chunk_id"]
        )

    def test_empty_and_invalid_windows(self):
        for size, overlap in [(0, 0), (100, 100), (100, -1)]:
            with self.assertRaises(ValueError):
                app.windows("d", "t", "a", size, overlap)
        with self.assertRaises(ValueError):
            app.windows("d", "t", " \n")
        self.assertEqual(len(app.windows("d", "t", "中" * 1000)), 1)

    def test_cache_bound_to_identity_and_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "batch.json"
            app.store_cache(p, "identity-a", {"value": 1})
            self.assertEqual(app.read_cache(p, "identity-a"), {"value": 1})
            with self.assertRaises(ValueError):
                app.read_cache(p, "identity-b")
            raw = json.loads(p.read_text(encoding="utf-8"))
            raw["data"]["value"] = 2
            p.write_text(json.dumps(raw), encoding="utf-8")
            with self.assertRaises(ValueError):
                app.read_cache(p, "identity-a")

    def test_cache_never_overwrites_existing(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "x.json"
            app.store_cache(p, "a", {"x": 1})
            with self.assertRaises(FileExistsError):
                app.store_cache(p, "a", {"x": 2})

    def test_manifest_digest_and_unique_count_are_required(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            p = root / "sources.jsonl"
            p.write_text('{"id":"a"}\n{"id":"b"}\n', encoding="utf-8")
            config = {"sources_sha256": app.sha(p), "documents": 2}
            self.assertEqual(len(app.read_sources(root, config)), 2)
            p.write_text('{"id":"a"}\n{"id":"a"}\n', encoding="utf-8")
            with self.assertRaises(ValueError):
                app.read_sources(root, config)
            config["sources_sha256"] = app.sha(p)
            with self.assertRaises(ValueError):
                app.read_sources(root, config)

    def test_metadata_is_repaired_on_resume_but_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            metadata = {"title": "synthetic", "source": "x"}
            app.ensure_metadata(root, metadata)
            (root / "document.json").unlink()
            app.ensure_metadata(root, metadata)
            self.assertEqual(app.load(root / "document.json"), metadata)
            with self.assertRaises(ValueError):
                app.ensure_metadata(root, {"different": 1})

    def test_rate_limit_retry_stops_before_next_request(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.bin"
            source.write_bytes(b"synthetic")
            entry = {
                "id": "sample",
                "title": "合成法",
                "source_path": str(source),
                "input_path": str(source),
                "source_sha256": app.sha(source),
                "input_sha256": app.sha(source),
            }
            config = {
                "size": 1000,
                "overlap": 200,
                "serializer": "test",
                "model": {"model": "test"},
            }
            calls = []

            def post(route, json):
                calls.append(json)
                (root / "STOP").touch()
                return SimpleNamespace(status_code=429)

            with (
                patch.object(app, "CONTEXT", (root, config)),
                patch.object(app, "HTTP", SimpleNamespace(post=post)),
                patch.object(
                    app,
                    "parse_document",
                    return_value=(
                        {"texts": []},
                        "第一条 正文",
                        {"needs_review": False},
                    ),
                ),
                patch.object(app.time, "sleep"),
            ):
                result = app.process(entry)
            self.assertEqual(result["status"], "interrupted")
            self.assertEqual(len(calls), 1)

    def test_vector_index_alignment_and_validation(self):
        rows = app.windows("d", "t", "中" * 1500)
        data = {
            "data": [
                {"index": 1, "embedding": [0, 1]},
                {"index": 0, "embedding": [1, 0]},
            ]
        }
        result = app.validate_vectors(data, rows, 2)
        self.assertEqual(result[0]["vector"], [1, 0])
        for invalid in [
            [],
            [{"index": 0, "embedding": [1, 0]}],
            [{"index": 0, "embedding": [1, 0]}, {"index": 0, "embedding": [1, 0]}],
        ]:
            with self.assertRaises(ValueError):
                app.validate_vectors({"data": invalid}, rows, 2)
        with self.assertRaises(ValueError):
            app.validate_vectors(
                {"data": [{"index": 0, "embedding": [float("nan"), 1]}]}, rows[:1], 2
            )

    def test_docling_plain_text_keeps_table_and_excludes_footer(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "synthetic.docx"
            doc = Document()
            doc.add_paragraph("合成法律")
            doc.add_paragraph("第一条 主文完整。")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "表格条件"
            table.cell(0, 1).text = "表格义务"
            doc.sections[0].footer.paragraphs[0].text = "－99－"
            doc.save(p)
            raw, text, quality = app.parse_document(p)
            self.assertIn("第一条 主文完整。", text)
            self.assertIn("表格条件", text)
            self.assertIn("表格义务", text)
            self.assertNotIn("－99－", text)
            self.assertTrue(raw["texts"])
            self.assertEqual(quality["conversion_status"], "success")

    def test_document_resume_uses_cache_and_tamper_is_rejected(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            doc = Document()
            doc.add_paragraph("第一条 合成正文。")
            doc.save(source)
            entry = {
                "id": "sample",
                "title": "合成法",
                "source_path": str(source),
                "input_path": str(source),
                "source_sha256": app.sha(source),
                "input_sha256": app.sha(source),
            }
            config = {
                "size": 1000,
                "overlap": 200,
                "serializer": "synthetic",
                "model": {"model": "test"},
            }
            calls = []

            def post(route, json):
                calls.append(json)
                data = {
                    "data": [
                        {"index": i, "embedding": [1.0] + [0.0] * 1023}
                        for i in range(len(json["input"]))
                    ],
                    "usage": {"total_tokens": 7},
                }
                return SimpleNamespace(
                    status_code=200, raise_for_status=lambda: None, json=lambda: data
                )

            with (
                patch.object(app, "CONTEXT", (root, config)),
                patch.object(app, "HTTP", SimpleNamespace(post=post)),
            ):
                first = app.process(entry)
                self.assertEqual(first["status"], "completed")
                second = app.process(entry)
                self.assertEqual(second, first)
                self.assertEqual(len(calls), 1)
                self.assertTrue(
                    calls[0]["input"][0].startswith("法律名称：合成法\n正文：")
                )
                chunks = root / "documents/sample/chunks.json"
                chunks.write_text("[]", encoding="utf-8")
                third = app.process(entry)
                self.assertEqual(third["status"], "failed")
                self.assertEqual(len(calls), 1)


if __name__ == "__main__":
    unittest.main()
